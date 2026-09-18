import csv
import json
from pathlib import Path

from app.rag.types import LoadedDocument


class TextLoader:
    def load(self, path: Path) -> list[LoadedDocument]:
        return [LoadedDocument(text=path.read_text(encoding="utf-8"))]


class JsonLoader:
    def load(self, path: Path) -> list[LoadedDocument]:
        value = json.loads(path.read_text(encoding="utf-8"))
        text = json.dumps(value, ensure_ascii=True, indent=2)
        return [LoadedDocument(text=text)]


class CsvLoader:
    def load(self, path: Path) -> list[LoadedDocument]:
        with path.open(newline="", encoding="utf-8") as file:
            rows = list(csv.DictReader(file))
        return [LoadedDocument(text=json.dumps(row, ensure_ascii=True)) for row in rows]


class PdfLoader:
    def load(self, path: Path) -> list[LoadedDocument]:
        from pypdf import PdfReader

        documents: list[LoadedDocument] = []
        for page_number, page in enumerate(PdfReader(str(path)).pages, start=1):
            documents.append(
                LoadedDocument(text=page.extract_text() or "", metadata={"page_number": page_number})
            )
        return documents


class DocxLoader:
    def load(self, path: Path) -> list[LoadedDocument]:
        from docx import Document

        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return [LoadedDocument(text=text)]
